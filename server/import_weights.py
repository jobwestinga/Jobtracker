"""Connect a local Weight_graph project to JobTracker through idempotent ops.

Usage: python -m server.import_weights --source /path/to/Weight_graph \
    --server https://your-server --token-file /path/to/sync_token

Use --dry-run to validate/count without changing anything. Spreadsheet support
is needed only on the importing computer, never in the server or desktop app.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import tempfile
import uuid
from datetime import datetime
from pathlib import Path

from jobtracker.services.weight_service import validate_measurement
from jobtracker.core.persistence.weight_codec import import_identity


def _record(moment, value, source):
    timestamp, weight, source = validate_measurement(moment.isoformat(), value, source)
    return {
        "uid": import_identity(timestamp, source),
        "measured_at": timestamp, "weight_kg": weight, "source": source,
    }


def load_scale(path: Path) -> list[dict]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    header = next((i for i, line in enumerate(lines) if line.startswith("Date;Time;Weight")), None)
    if header is None:
        raise ValueError(f"Missing HealthManager table in {path.name}")
    records = {}
    for number, row in enumerate(csv.DictReader(io.StringIO("\n".join(lines[header:])), delimiter=";"), header + 2):
        value = (row.get("Weight") or "").strip()
        if not value:
            continue
        try:
            moment = datetime.strptime(f"{row['Date']} {row['Time']}", "%d-%m-%Y %H:%M")
            record = _record(moment, value, "scale")
        except (ValueError, KeyError) as exc:
            raise ValueError(f"Invalid measurement in {path.name}, row {number}") from exc
        records[record["uid"]] = record
    return sorted(records.values(), key=lambda row: row["measured_at"])


def load_workbook(path: Path, source: str) -> list[dict]:
    if not path.exists():
        return []
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("Install requirements-weight-import.txt to read the source workbooks") from None
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = {}
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headings = list(next(rows, []))
        if "Date" not in headings or "Weight" not in headings:
            raise ValueError(f"{path.name} needs Date and Weight columns")
        date_index, weight_index = headings.index("Date"), headings.index("Weight")
        for number, row in enumerate(rows, 2):
            moment, value = row[date_index], row[weight_index]
            # Matches Weight_graph's dropna: blank source cells are not readings.
            if moment is None or value is None:
                continue
            try:
                if isinstance(moment, str):
                    moment = datetime.fromisoformat(moment)
                elif not isinstance(moment, datetime):
                    moment = datetime.combine(moment, datetime.min.time())
                record = _record(moment, value, source)
            except (ValueError, TypeError, AttributeError) as exc:
                raise ValueError(f"Invalid measurement in {path.name}, row {number}") from exc
            records[record["uid"]] = record
    finally:
        workbook.close()
    return sorted(records.values(), key=lambda row: row["measured_at"])


def load_project(source: Path) -> list[dict]:
    data = source / "data" if (source / "data").is_dir() else source
    exports = [p for p in data.iterdir() if p.suffix.lower() == ".csv" and any(
        line.startswith("Date;Time;Weight") for line in p.read_text(encoding="utf-8-sig").splitlines()
    )]
    scale = load_scale(max(exports, key=lambda p: p.stat().st_mtime)) if exports else []
    manual = load_workbook(data / "WeightData.xlsx", "manual")
    historic = load_workbook(data / "HistoricWeight.xlsx", "historic")
    # Same calendar-day precedence as the existing plotting project.
    occupied = {row["measured_at"][:10] for row in scale}
    merged = scale + [row for row in manual if row["measured_at"][:10] not in occupied]
    occupied = {row["measured_at"][:10] for row in merged}
    merged.extend(row for row in historic if row["measured_at"][:10] not in occupied)
    if not merged:
        raise ValueError("No weight measurements found in the source project")
    return sorted(merged, key=lambda row: row["measured_at"])


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--server")
    parser.add_argument("--token-file", type=Path)
    parser.add_argument("--database", type=Path, help="Import into a local database instead of the server")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--export-json", type=Path, help="Write merged readings for the existing plotting project")
    args = parser.parse_args(argv)
    if args.export_json:
        if not args.server or not args.token_file:
            parser.error("Export requires --server and --token-file")
        from jobtracker.sync.client import SyncClient
        client = SyncClient(args.server, args.token_file.read_text().strip())
        payload = client._request("GET", "/api/weight")
        args.export_json.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary = tempfile.mkstemp(prefix=".jobtracker-weights-", dir=args.export_json.parent)
        try:
            with os.fdopen(descriptor, "w") as output:
                json.dump(payload["measurements"], output, indent=2)
                output.flush()
                os.fsync(output.fileno())
            os.replace(temporary, args.export_json)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
        print(f"Exported {len(payload['measurements'])} measurements")
        return
    measurements = load_project(args.source)
    print(f"Validated {len(measurements)} merged measurements")
    if args.dry_run:
        return
    if args.database:
        os.environ["JOBTRACKER_DB_PATH"] = str(args.database)
        from jobtracker.core.database import Database
        from jobtracker.services.weight_service import WeightService
        database = Database(args.database)
        try:
            print(WeightService(database).import_rows(measurements))
        finally:
            database.connection.close()
    elif args.server and args.token_file:
        from jobtracker.sync.client import SyncClient
        client = SyncClient(args.server, args.token_file.read_text().strip())
        for offset in range(0, len(measurements), 2000):
            batch = measurements[offset:offset + 2000]
            # Receipt identity is per payload, so updating the source adds a new
            # receipt; row identities still make repeated imports harmless.
            identity = str(uuid.uuid5(uuid.NAMESPACE_URL, json.dumps(batch, sort_keys=True)))
            response = client.send_ops([{"op_id": identity, "op": "import_weights", "params": {"measurements": batch}}])
            print(response["applied"][0]["result"])
    else:
        parser.error("Use --dry-run, --database, or both --server and --token-file")


if __name__ == "__main__":
    main()
